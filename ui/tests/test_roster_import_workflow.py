from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from models.build_model import BuildRoster, PlayerBuild
from models.roster_model import RosterMember
from services.build_service import BuildService
from services.eso_database import EsoDatabase
from services.roster_service import RosterService
from ui.roster_import_workflow import (
    ImportedBuildCandidate,
    ImportedRosterMember,
    RosterImportParser,
    RosterImportPlan,
    apply_roster_import,
)


def _save(workbook: Workbook, path: Path) -> Path:
    workbook.save(path)
    return path


def test_hh_style_workbook_parses_personal_and_shared_builds(tmp_path: Path) -> None:
    workbook = Workbook()
    info = workbook.active
    info.title = "Info"
    info["A15"] = "Pilly - Den"
    info["B15"] = "Jarakeen"
    info["A16"] = "Pure Arc"
    info["B16"] = "Pippin"
    info["A17"] = "Pure Arc"
    info["B17"] = "Fu"

    healer = workbook.create_sheet("Jarakeen - Den Heroism Healer")
    healer["A1"] = "Jarakeen"
    healer["A2"] = "Den Healer Major Heroism"
    healer["A3"] = "Pure Den - Trash"
    for column, value in enumerate(("Piece", "Set", "Weight", "Trait", "Enchant"), start=1):
        healer.cell(4, column, value)
    gear_rows = (
        ("Head", "Ozezan", "Light", "Divines", "Max Mag"),
        ("Shoulders", "Ozezan", "Light", "Divines", "Max Mag"),
        ("Chest", "SPC", "Light", "Divines", "Max Mag"),
        ("Necklace", "Powerful Assault", "Jewelry", "Swift", "Mag Cost Redux"),
        ("Ring", "Powerful Assault", "Jewelry", "Infused", "Spell Dmg"),
        ("Ring", "Powerful Assault", "Jewelry", "Infused", "Spell Dmg"),
        ("Weapon1", "Perfected Master's", "Resto Staff", "Powered", "Weap Dmg"),
        ("Weapon2", "Powerful Assault", "Frost Staff", "Infused", "Crusher"),
    )
    for row, values in enumerate(gear_rows, start=5):
        for column, value in enumerate(values, start=1):
            healer.cell(row, column, value)
    healer.cell(13, 1, "Resto Bar")
    healer.cell(13, 2, "Frost Bar")
    healer.cell(13, 3, "Warfare CP")
    healer.cell(13, 4, "Fitness CP")
    healer.cell(13, 5, "Craft CP")
    for index, (front, back) in enumerate(
        (
            ("Combat Prayer", "Deceptive Predator"),
            ("Energy Orb", "Expansive Frost Cloak"),
            ("Illustrious Healing", "Echoing Vigor"),
            ("Radiating Regen", "Elemental Blockade"),
            ("Budding Seeds", "Overflowing Altar"),
            ("Reviving Barrier", "Aggressive Horn"),
        ),
        start=14,
    ):
        healer.cell(index, 1, front)
        healer.cell(index, 2, back)
    healer.cell(14, 3, "From the Brink")
    healer.cell(14, 4, "Boundless Vitality")
    healer.cell(18, 5, "Attributes: 64 Mag")
    healer.cell(19, 5, "Food: Clockwork Citrus Filet")
    healer.cell(20, 5, "Mundus: Ritual")
    healer.cell(21, 5, "Potions: Tri-stat")

    shared = workbook.create_sheet("Pippin and Fu Arc DD")
    shared["A1"] = "Pippin, Fu"
    shared["A2"] = "Arc DD"
    shared["A3"] = "Pure Arc - Trash"
    for column, value in enumerate(("Piece", "Set", "Weight", "Trait", "Enchant"), start=1):
        shared.cell(4, column, value)
    shared.append(["Head", "Slimecraw", "Medium", "Divines", "Max Stam"])
    shared.append(["Weapon1", "Null Arca", "Dagger", "Charged", "Fire Dmg"])
    shared.append(["Weapon1", "Null Arca", "Dagger", "Charged", "Poison Dmg"])
    shared.append(["Weapon2", "Maelstrom Perfected", "Greatsword", "Infused", "Weap Dmg"])
    shared.append(["Daggers Bar", "Greatsword Bar", "Warfare CP", "Fitness CP", "Craft CP"])
    for front, back in (
        ("Cephaliarch's Flail", "Stampede"),
        ("Barbed Trap", "Carve"),
        ("Pragmatic Fatecarver", "Fulminating Rune"),
        ("Quick Cloak", "Scalding Rune"),
        ("Camo Hunter", "Inspired Scholarship"),
        ("Flawless Dawnbreaker", "Languid Eye"),
    ):
        shared.append([front, back])

    path = _save(workbook, tmp_path / "Swine and Punishment HH U50.xlsx")
    plan = RosterImportParser().parse_file(path)

    by_tag = {member.gamertag: member for member in plan.members}
    jara = by_tag["Jarakeen"]
    assert jara.eso_class == "Warden"
    assert jara.primary_role == "Healer"
    assert len(jara.builds) == 1
    assert jara.builds[0].payload["Armor"]["Head"]["Set"] == "Ozezan"
    assert jara.builds[0].payload["FrontBarWeapon"]["WeaponType"] == "Restoration Staff"
    assert jara.builds[0].payload["BackBarWeapon"]["WeaponType"] == "Ice Staff"
    assert jara.builds[0].payload["FrontBarSkills"][0] == "Combat Prayer"

    assert by_tag["Pippin"].builds
    assert by_tag["Fu"].builds
    assert by_tag["Pippin"].builds[0].payload["FrontBarOffHand"]["WeaponType"] == "Dagger"


def test_sectioned_gs_workbook_attaches_templates_by_role_and_class(tmp_path: Path) -> None:
    workbook = Workbook()
    roster = workbook.active
    roster.title = "TEAM ROSTER"
    roster["B7"] = "TANKS"
    roster["B8"] = "Wicked - Main - Minor Courage"
    roster["E8"] = "Nightblade - Cutthroats"
    roster["F8"] = "Yoln / Xoryn"
    roster["C11"] = "HEALERS"
    roster["B12"] = "left"
    roster["C12"] = "Sparkle - Major Force"
    roster["E12"] = "Arcanist - Ink-scribes Verb"
    roster["B13"] = "right"
    roster["C13"] = "Jarakeen - Major Brittle, Major Heroism"
    roster["E13"] = "Warden - Tundra's Maw"
    roster["C15"] = "BUFF DPS"
    roster["B17"] = "M3 / 6 (CA)"
    roster["C17"] = "Feral Puppers - Major Berserk, Zen's"
    roster["E17"] = "Dragon Knight - LFTF"
    roster["C21"] = "FULL DAMY DD"
    roster["B22"] = "M1 / 2 (CA)"
    roster["C22"] = "Kye - Portal"
    roster["E22"] = "Templar - Bright Harbinger"

    den = workbook.create_sheet("DEN HEALER")
    den["B2"] = "GEAR"
    den["D2"] = "Serpent's Disdain body"
    den["B3"] = "PURE CLASS"
    den["D3"] = "WARDEN"
    den["B4"] = "BLUE CP"
    den["D4"] = "From the Brink"
    den["B6"] = "FRONT BAR SKILLS"
    den["D6"] = "Combat Prayer, Energy Orb, Budding Seeds, Radiating Regen, Enchanted Growth, Barrier"
    den["B7"] = "BACK BAR SKILLS"
    den["D7"] = "Elemental Blockade, Frost Cloak, Vigor, Altar, Fetcher Infection, Aggressive Horn"

    plar = workbook.create_sheet("TEMPLAR DD")
    plar["B2"] = "GEAR"
    plar["D2"] = "Lancer body"
    plar["B3"] = "PURE CLASS"
    plar["D3"] = "TEMPLAR"
    plar["B6"] = "FRONT BAR SKILLS"
    plar["D6"] = "Biting Jabs, Barbed Trap, Power of the Light, Blazing Spear, Radiant Oppression, Dawnbreaker"
    plar["B7"] = "BACK BAR SKILLS"
    plar["D7"] = "Stampede, Carve, Solar Barrage, Vampire's Bane, Ritual of Retribution, Onslaught"

    path = _save(workbook, tmp_path / "U50 GS comp.xlsx")
    plan = RosterImportParser().parse_file(path)
    by_tag = {member.gamertag: member for member in plan.members}

    assert by_tag["Jarakeen"].primary_role == "Healer"
    assert by_tag["Jarakeen"].eso_class == "Warden"
    assert [build.build_name for build in by_tag["Jarakeen"].builds] == ["DEN HEALER"]
    assert by_tag["Jarakeen"].builds[0].payload["FrontBarSkills"][0] == "Combat Prayer"
    assert [build.build_name for build in by_tag["Kye"].builds] == ["TEMPLAR DD"]


def test_apply_import_preserves_unrelated_roster_and_builds_and_assigns_team(tmp_path: Path) -> None:
    roster_service = RosterService(EsoDatabase(tmp_path / "eso.db"))
    roster_service.create_member(
        RosterMember(
            PlayerName="Existing Player",
            CharacterName="Existing Character",
            EsoClass="Sorcerer",
            PrimaryRole="Damage Dealer",
            Team="Old Team",
        )
    )

    build_service = BuildService(tmp_path / "builds.json")
    build_service.save(
        BuildRoster(
            Members=[
                PlayerBuild(
                    Name="Existing Character",
                    Gamertag="Existing Player",
                    BuildName="Old Build",
                    EsoClass="Sorcerer",
                    Role="Damage Dealer",
                    Food="Old Food",
                )
            ]
        )
    )

    imported = ImportedRosterMember(
        gamertag="Jarakeen",
        character_name="Magrat",
        eso_class="Warden",
        primary_role="Healer",
        assignment="Pillager / Heroism",
        builds=[
            ImportedBuildCandidate(
                gamertag="Jarakeen",
                build_name="HH U50 Healer",
                eso_class="Warden",
                role="Healer",
                assignment="Pillager / Heroism",
                payload={
                    "FrontBarSkills": ["Combat Prayer", "Energy Orb", "", "", "", ""],
                    "BackBarSkills": ["Elemental Blockade", "", "", "", "", "Aggressive Horn"],
                },
            )
        ],
    )
    plan = RosterImportPlan(tmp_path / "source.xlsx", "Swine and Punishment", [imported])

    result = apply_roster_import(plan, roster_service, build_service, import_builds=True)

    assert result.created_roster_members == 1
    assert result.imported_builds == 1
    assert {member.PlayerName for member in roster_service.list_members()} == {
        "Existing Player",
        "Jarakeen",
    }

    saved = build_service.load().Members
    saved_keys = {(build.Gamertag, build.Name, build.BuildName) for build in saved}
    assert ("Existing Player", "Existing Character", "Old Build") in saved_keys
    assert ("Jarakeen", "Magrat", "HH U50 Healer") in saved_keys

    catalog = build_service.canonical.catalog_service.load()
    imported_build = next(
        build for build in catalog["builds"] if build.get("name") == "HH U50 Healer"
    )
    assignments = build_service.canonical.catalog_service.assignments_for_build(
        imported_build["build_id"]
    )
    assert assignments[0]["team_name"] == "Swine and Punishment"
    assert assignments[0]["raid_role"] == "Healer"
    assert assignments[0]["slot_name"] == "Pillager / Heroism"
