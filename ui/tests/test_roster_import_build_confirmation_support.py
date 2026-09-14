from __future__ import annotations

from types import SimpleNamespace

from openpyxl import Workbook

from ui import roster_import_build_confirmation_support as support
from ui.roster_import_workflow import ImportedBuildCandidate


def test_scribed_recipe_rows_are_preserved_for_confirmation() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(17, 1).value = "Front Bar"
    sheet.cell(17, 2).value = "Back Bar"
    sheet.cell(17, 3).value = "Warfare CP"
    sheet.cell(22, 3).value = "Scribed Skills"
    sheet.cell(23, 3).value = "Banner: Shock Dmg, Class Flourish, Courage"

    candidate = ImportedBuildCandidate(
        gamertag="Player",
        build_name="Test Build",
        eso_class="Necromancer",
        payload={
            "FrontBarSkills": ["Banner", "", "", "", "", ""],
            "BackBarSkills": ["", "", "", "", "", ""],
        },
    )
    parser = SimpleNamespace(_bar_header_row=lambda _sheet, _start_col: 17)

    support._read_scribed_recipes(parser, sheet, 1, candidate)

    recipes = candidate.payload["ScribedSkillRecipes"]
    assert recipes == [
        {
            "ResultName": "Banner",
            "Grimoire": "",
            "Focus": "Shock Damage",
            "Signature": "Class Flourish",
            "Affix": "Courage",
        }
    ]
    assert candidate.payload["ScribedSkills"] == ["Banner"]


def test_ambiguous_potion_options_are_not_promoted_to_one_selection() -> None:
    candidate = ImportedBuildCandidate(
        gamertag="Player",
        build_name="Test Build",
        payload={"Potion": "Tri-stat, Bi-stat Stam", "Notes": "Imported"},
    )

    support._preserve_ambiguous_potion(candidate)

    assert candidate.payload["Potion"] == ""
    assert candidate.payload["_ImportPotionOptions"] == ["Tri-stat", "Bi-stat Stam"]
    assert "Imported potion options (choose one): Tri-stat, Bi-stat Stam" in candidate.payload["Notes"]


def test_build_check_requires_potion_and_scribed_grimoire_confirmation(monkeypatch) -> None:
    candidate = ImportedBuildCandidate(
        gamertag="Player",
        build_name="Test Build",
        eso_class="Necromancer",
        payload={
            "Potion": "",
            "_ImportPotionOptions": ["Tri-stat", "Bi-stat Stam"],
            "FrontBarSkills": ["Banner", "", "", "", "", ""],
            "BackBarSkills": ["", "", "", "", "", ""],
            "ScribedSkillRecipes": [
                {
                    "ResultName": "Banner",
                    "Grimoire": "",
                    "Focus": "Shock Damage",
                    "Signature": "Class Flourish",
                    "Affix": "Courage",
                }
            ],
        },
    )
    monkeypatch.setattr(support, "_ability_is_crafted", lambda *_args: True)

    issues = support._candidate_issues(candidate)

    assert "choose potion: Tri-stat / Bi-stat Stam" in issues
    assert "confirm Banner: Grimoire" in issues


def test_single_unresolved_potion_label_is_still_flagged(monkeypatch) -> None:
    candidate = ImportedBuildCandidate(
        gamertag="Player",
        build_name="Test Build",
        payload={"Potion": "Tri-stat", "FrontBarSkills": [], "BackBarSkills": []},
    )
    monkeypatch.setattr(support, "_potion_resolves", lambda _label: False)

    issues = support._candidate_issues(candidate)

    assert issues == ("confirm canonical potion label: Tri-stat",)
