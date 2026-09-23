from __future__ import annotations

"""Import external raid rosters into the existing Player -> Character -> Build model.

The workbook parser is deliberately heuristic because raid-lead sheets are
human documents, not database exports.  Heuristics stop at the preview boundary:
unknown character identity stays unresolved, and no roster/build state is
written until the user reviews the detected rows.
"""

import csv
import json
import re
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QCheckBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
)

from engine.config import get_data_dir
from models.build_model import BuildRoster, PlayerBuild
from models.roster_model import ESO_CLASSES, RosterMember
from services.build_service import BuildService
from services.user_safety_snapshot_service import UserSafetySnapshotService


_INSTALLED = False


@dataclass
class ImportedBuildCandidate:
    gamertag: str
    build_name: str
    eso_class: str = ""
    role: str = ""
    assignment: str = ""
    source_sheet: str = ""
    payload: dict[str, Any] = field(default_factory=dict)


@dataclass
class ImportedRosterMember:
    gamertag: str
    character_name: str = ""
    eso_class: str = ""
    primary_role: str = ""
    secondary_role: str = ""
    status: str = "Active"
    assignment: str = ""
    source: str = ""
    builds: list[ImportedBuildCandidate] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    selected: bool = True


@dataclass
class RosterImportPlan:
    source_path: Path
    team_name: str
    members: list[ImportedRosterMember] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    source_kind: str = ""

    @property
    def build_count(self) -> int:
        return sum(len(member.builds) for member in self.members)


@dataclass(frozen=True)
class RosterImportResult:
    created_roster_members: int = 0
    updated_roster_members: int = 0
    imported_builds: int = 0
    skipped_builds: int = 0
    warnings: tuple[str, ...] = ()


_COLUMN_ALIASES = {
    "player": "gamertag",
    "player name": "gamertag",
    "gamertag": "gamertag",
    "gamer tag": "gamertag",
    "account": "gamertag",
    "@name": "gamertag",
    "character": "character_name",
    "character name": "character_name",
    "toon": "character_name",
    "class": "eso_class",
    "eso class": "eso_class",
    "role": "primary_role",
    "primary role": "primary_role",
    "secondary role": "secondary_role",
    "off role": "secondary_role",
    "status": "status",
    "assignment": "assignment",
    "notes": "assignment",
    "team": "team",
    "group": "team",
}


def _text(value: Any) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def _class_from_text(value: Any) -> str:
    raw = f" {_text(value).casefold()} "
    aliases = (
        (("dragonknight", "dragon knight", " dk "), "Dragonknight"),
        (("sorcerer", " sorc"), "Sorcerer"),
        (("nightblade", " nb "), "Nightblade"),
        (("templar", " plar"), "Templar"),
        (("warden", " den "), "Warden"),
        (("necromancer", " necro", " cro "), "Necromancer"),
        (("arcanist", " arc "), "Arcanist"),
    )
    for tokens, result in aliases:
        if any(token in raw for token in tokens):
            return result
    return ""


def _role_from_text(value: Any, *, default_damage: bool = False) -> str:
    raw = f" {_text(value).casefold()} "
    if "healer" in raw or " heal " in raw or "mapa" in raw or "pilly" in raw:
        return "Healer"
    if "tank" in raw or " mt " in raw or " ot " in raw or raw.strip().startswith(("mt", "ot")):
        return "Tank"
    if any(token in raw for token in ("support dd", "support dps", "zens", "z'en", "brittle")):
        return "DD"
    if any(token in raw for token in (" dps", " dd ", "werewolf", " ww ", "banner", "knife", "force", "colorless")):
        return "DD"
    return "DD" if default_damage else ""


def _normalize_role(value: Any) -> str:
    raw = _text(value)
    folded = raw.casefold()
    if folded in {"dd", "dps", "damage", "damage dealer"}:
        return "DD"
    if folded in {"support dd", "support dps", "support damage", "support damage dealer"}:
        return "DD"
    if folded in {"heal", "heals", "healer"}:
        return "Healer"
    if folded in {"tank", "mt", "ot", "main tank", "off tank"}:
        return "Tank"
    return raw


def _normalize_class(value: Any) -> str:
    raw = _text(value)
    if raw in ESO_CLASSES:
        return raw
    return _class_from_text(raw) or raw


def _weapon_type(value: Any) -> str:
    raw = _text(value)
    folded = raw.casefold()
    if "resto" in folded:
        return "Restoration Staff"
    if "frost" in folded or "ice staff" in folded:
        return "Ice Staff"
    if "inferno" in folded or "fire staff" in folded:
        return "Inferno Staff"
    if "lightning" in folded:
        return "Lightning Staff"
    if "bow" in folded:
        return "Bow"
    if "2h" in folded or "two handed" in folded or "two-handed" in folded or "greatsword" in folded:
        return "Two-Handed"
    if "dual" in folded:
        return "Dual Wield"
    if "sword" in folded and "shield" in folded:
        return "One Hand and Shield"
    for one_hand in ("dagger", "sword", "axe", "mace", "shield"):
        if one_hand in folded:
            return one_hand.title()
    return raw


def _split_skills(value: Any) -> list[str]:
    raw = _text(value)
    if not raw:
        return [""] * 6
    parts = [
        part.strip()
        for part in re.split(r"\s*(?:,|\||\s/\s)\s*", raw)
        if part.strip()
    ]
    return (parts + [""] * 6)[:6]


def _slot_payload(set_name: Any, weight: Any, trait: Any, enchant: Any, *, weapon: bool = False) -> dict[str, str]:
    result = {
        "Set": _text(set_name),
        "Trait": _text(trait),
        "Enchant": _text(enchant),
        "Weight": _text(weight),
    }
    if weapon:
        result["WeaponType"] = _weapon_type(weight)
    return result


def _empty_armor_payload() -> dict[str, dict[str, str]]:
    return {
        slot: {
            "Set": "",
            "Set2": "",
            "Quality": "",
            "Trait": "",
            "Enchant": "",
            "EnchantTier": "",
            "Level": "",
            "Weight": "",
        }
        for slot in ("Head", "Shoulders", "Chest", "Hands", "Waist", "Legs", "Feet")
    }


class RosterImportParser:
    """Read Excel/CSV/JSON roster sources into a mutation-free preview plan."""

    def parse_file(self, path: Path) -> RosterImportPlan:
        path = Path(path)
        suffix = path.suffix.casefold()
        if suffix in {".xlsx", ".xlsm"}:
            return self._parse_workbook(path)
        if suffix == ".csv":
            return self._parse_csv(path)
        if suffix == ".json":
            return self._parse_json(path)
        raise ValueError("Roster import supports .xlsx, .xlsm, .csv, and .json files.")

    def _parse_workbook(self, path: Path) -> RosterImportPlan:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Roster workbook import needs openpyxl.") from exc

        workbook = load_workbook(path, data_only=True)
        plan = RosterImportPlan(path, path.stem.strip(), source_kind="workbook")
        if "TEAM ROSTER" in workbook.sheetnames:
            self._parse_sectioned_team_roster(workbook["TEAM ROSTER"], plan)
            self._attach_role_templates(workbook, plan)
        else:
            if "Info" in workbook.sheetnames:
                self._parse_info_assignments(workbook["Info"], plan)
            self._parse_personal_build_sheets(workbook, plan)
        if not plan.members:
            plan.warnings.append("No roster members were detected in this workbook.")
        return plan

    def _parse_csv(self, path: Path) -> RosterImportPlan:
        plan = RosterImportPlan(path, path.stem.strip(), source_kind="csv")
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            for raw_row in reader:
                normalized: dict[str, str] = {}
                for key, value in (raw_row or {}).items():
                    canonical = _COLUMN_ALIASES.get(_text(key).casefold())
                    if canonical:
                        normalized[canonical] = _text(value)
                gamertag = normalized.get("gamertag", "")
                if not gamertag:
                    continue
                if normalized.get("team") and not plan.members:
                    plan.team_name = normalized["team"]
                plan.members.append(
                    ImportedRosterMember(
                        gamertag=gamertag,
                        character_name=normalized.get("character_name", ""),
                        eso_class=_normalize_class(normalized.get("eso_class", "")),
                        primary_role=_normalize_role(normalized.get("primary_role", "")),
                        secondary_role=_normalize_role(normalized.get("secondary_role", "")),
                        status=normalized.get("status", "") or "Active",
                        assignment=normalized.get("assignment", ""),
                        source=path.name,
                    )
                )
        return plan

    def _parse_json(self, path: Path) -> RosterImportPlan:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, list):
            rows = payload
            team_name = path.stem.strip()
        elif isinstance(payload, dict):
            rows = payload.get("Members") or payload.get("members") or payload.get("roster") or []
            team_name = _text(payload.get("TeamName") or payload.get("team_name") or path.stem)
        else:
            raise ValueError("Roster JSON must contain a list or an object with Members/members/roster.")
        plan = RosterImportPlan(path, team_name, source_kind="json")
        for raw in rows:
            if not isinstance(raw, dict):
                continue
            gamertag = _text(raw.get("Gamertag") or raw.get("PlayerName") or raw.get("player") or raw.get("gamertag"))
            if not gamertag:
                continue
            member = ImportedRosterMember(
                gamertag=gamertag,
                character_name=_text(raw.get("CharacterName") or raw.get("Name") or raw.get("character")),
                eso_class=_normalize_class(raw.get("EsoClass") or raw.get("Class") or raw.get("class")),
                primary_role=_normalize_role(raw.get("PrimaryRole") or raw.get("Role") or raw.get("role")),
                secondary_role=_normalize_role(raw.get("SecondaryRole") or raw.get("secondary_role")),
                status=_text(raw.get("Status") or raw.get("status")) or "Active",
                assignment=_text(raw.get("Assignment") or raw.get("Notes") or raw.get("assignment")),
                source=path.name,
            )
            raw_builds = raw.get("Builds") or raw.get("builds") or []
            if isinstance(raw_builds, dict):
                raw_builds = [raw_builds]
            for index, raw_build in enumerate(raw_builds):
                if not isinstance(raw_build, dict):
                    continue
                build_payload = deepcopy(raw_build)
                build_name = _text(build_payload.get("BuildName") or build_payload.get("name")) or f"Imported {index + 1}"
                member.builds.append(
                    ImportedBuildCandidate(
                        gamertag=gamertag,
                        build_name=build_name,
                        eso_class=member.eso_class,
                        role=member.primary_role,
                        assignment=member.assignment,
                        source_sheet=path.name,
                        payload=build_payload,
                    )
                )
            plan.members.append(member)
        return plan

    @staticmethod
    def _member_by_tag(plan: RosterImportPlan, gamertag: str) -> ImportedRosterMember | None:
        wanted = _text(gamertag).casefold()
        return next((member for member in plan.members if member.gamertag.casefold() == wanted), None)

    def _ensure_member(
        self,
        plan: RosterImportPlan,
        *,
        gamertag: str,
        eso_class: str = "",
        role: str = "",
        assignment: str = "",
        source: str = "",
        authoritative_role: bool = False,
    ) -> ImportedRosterMember:
        gamertag = _text(gamertag)
        existing = self._member_by_tag(plan, gamertag)
        if existing is None:
            existing = ImportedRosterMember(
                gamertag=gamertag,
                eso_class=_normalize_class(eso_class),
                primary_role=_normalize_role(role),
                assignment=_text(assignment),
                source=source,
            )
            plan.members.append(existing)
            return existing
        if eso_class and (authoritative_role or not existing.eso_class):
            existing.eso_class = _normalize_class(eso_class)
        if role and (authoritative_role or not existing.primary_role):
            existing.primary_role = _normalize_role(role)
        if assignment and assignment.casefold() not in existing.assignment.casefold():
            existing.assignment = " · ".join(piece for piece in (existing.assignment, _text(assignment)) if piece)
        if source:
            existing.source = existing.source or source
        return existing

    def _parse_info_assignments(self, sheet, plan: RosterImportPlan) -> None:
        for row in range(1, min(sheet.max_row, 80) + 1):
            label = _text(sheet.cell(row, 1).value)
            player = _text(sheet.cell(row, 2).value)
            if not label or not player:
                continue
            eso_class = _class_from_text(label)
            role = _role_from_text(label, default_damage=bool(eso_class))
            if not eso_class and not role:
                continue
            self._ensure_member(
                plan,
                gamertag=player,
                eso_class=eso_class,
                role=role,
                assignment=label,
                source=sheet.title,
            )

        for column, label in ((1, "Left Stack"), (2, "Right Stack")):
            for row in range(4, min(sheet.max_row, 20) + 1):
                player = _text(sheet.cell(row, column).value)
                if not player or player.casefold() in {"left stack", "right stack"}:
                    continue
                member = self._member_by_tag(plan, player)
                if member and label.casefold() not in member.assignment.casefold():
                    member.assignment = " · ".join(piece for piece in (member.assignment, label) if piece)

    def _parse_personal_build_sheets(self, workbook, plan: RosterImportPlan) -> None:
        for sheet in workbook.worksheets:
            if sheet.title.casefold() in {"info", "team roster"}:
                continue
            if _text(sheet.cell(4, 1).value).casefold() != "piece":
                continue
            raw_players = _text(sheet.cell(1, 1).value)
            role_text = _text(sheet.cell(2, 1).value)
            if not raw_players or not role_text:
                continue
            gamertags = [piece.strip() for piece in raw_players.split(",") if piece.strip()]
            eso_class = _class_from_text(role_text)
            role = _role_from_text(role_text, default_damage=True)
            loadouts = [
                candidate
                for start_col in range(1, sheet.max_column + 1)
                if _text(sheet.cell(4, start_col).value).casefold() == "piece"
                if (
                    candidate := self._parse_personal_loadout(
                        sheet,
                        start_col=start_col,
                        gamertag="",
                        role_text=role_text,
                        eso_class=eso_class,
                        role=role,
                    )
                ) is not None
            ]
            for gamertag in gamertags:
                member = self._ensure_member(
                    plan,
                    gamertag=gamertag,
                    eso_class=eso_class,
                    role=role,
                    assignment=role_text,
                    source=sheet.title,
                    authoritative_role=True,
                )
                for template in loadouts:
                    candidate = deepcopy(template)
                    candidate.gamertag = gamertag
                    candidate.payload["Gamertag"] = gamertag
                    member.builds.append(candidate)

    @staticmethod
    def _bar_header_row(sheet, start_col: int) -> int | None:
        for row in range(5, min(sheet.max_row, 45) + 1):
            left = _text(sheet.cell(row, start_col).value).casefold()
            right = _text(sheet.cell(row, start_col + 1).value).casefold()
            if "bar" in left and "bar" in right:
                return row
        return None

    def _parse_personal_loadout(
        self,
        sheet,
        *,
        start_col: int,
        gamertag: str,
        role_text: str,
        eso_class: str,
        role: str,
    ) -> ImportedBuildCandidate | None:
        build_name = _text(sheet.cell(3, start_col).value) or _text(sheet.title)
        bar_row = self._bar_header_row(sheet, start_col)
        gear_end = (bar_row - 1) if bar_row else min(sheet.max_row, 24)
        armor = _empty_armor_payload()
        payload: dict[str, Any] = {
            "Gamertag": gamertag,
            "BuildName": build_name,
            "EsoClass": eso_class,
            "Role": role,
            "Armor": armor,
            "FrontBarSkills": [""] * 6,
            "BackBarSkills": [""] * 6,
            "ChampionPoints": [],
            "Werewolf": "ww" in role_text.casefold() or "werewolf" in role_text.casefold(),
            "Notes": f"Imported from {sheet.title}: {role_text}",
        }
        ring_index = 0
        weapon1_count = 0
        weapon2_count = 0
        meaningful_gear = False

        for row in range(5, gear_end + 1):
            piece = _text(sheet.cell(row, start_col).value)
            if not piece:
                continue
            set_name = sheet.cell(row, start_col + 1).value
            weight = sheet.cell(row, start_col + 2).value
            trait = sheet.cell(row, start_col + 3).value
            enchant = sheet.cell(row, start_col + 4).value
            key = piece.casefold()
            slot = {
                "head": "Head",
                "shoulders": "Shoulders",
                "chest": "Chest",
                "arms": "Hands",
                "hands": "Hands",
                "gloves": "Hands",
                "waist": "Waist",
                "legs": "Legs",
                "boots": "Feet",
                "feet": "Feet",
            }.get(key)
            if slot:
                armor[slot].update(_slot_payload(set_name, weight, trait, enchant))
                meaningful_gear = meaningful_gear or bool(_text(set_name))
            elif key == "necklace":
                payload["Necklace"] = _slot_payload(set_name, weight, trait, enchant)
                meaningful_gear = meaningful_gear or bool(_text(set_name))
            elif key == "ring":
                ring_index += 1
                payload["Ring1" if ring_index == 1 else "Ring2"] = _slot_payload(set_name, weight, trait, enchant)
                meaningful_gear = meaningful_gear or bool(_text(set_name))
            elif key in {"weapon1", "front bar weapon", "front weapon"}:
                weapon1_count += 1
                payload["FrontBarWeapon" if weapon1_count == 1 else "FrontBarOffHand"] = _slot_payload(
                    set_name, weight, trait, enchant, weapon=True
                )
                meaningful_gear = meaningful_gear or bool(_text(set_name))
            elif key in {"weapon2", "back bar weapon", "back weapon"}:
                weapon2_count += 1
                payload["BackBarWeapon" if weapon2_count == 1 else "BackBarOffHand"] = _slot_payload(
                    set_name, weight, trait, enchant, weapon=True
                )
                meaningful_gear = meaningful_gear or bool(_text(set_name))

        if bar_row:
            payload["FrontBarSkills"] = [
                _text(sheet.cell(row, start_col).value)
                for row in range(bar_row + 1, min(bar_row + 7, sheet.max_row + 1))
            ]
            payload["BackBarSkills"] = [
                _text(sheet.cell(row, start_col + 1).value)
                for row in range(bar_row + 1, min(bar_row + 7, sheet.max_row + 1))
            ]
            payload["FrontBarSkills"] = (payload["FrontBarSkills"] + [""] * 6)[:6]
            payload["BackBarSkills"] = (payload["BackBarSkills"] + [""] * 6)[:6]

            champion_points: list[dict[str, str]] = []
            for column in range(start_col + 2, min(start_col + 5, sheet.max_column + 1)):
                if "cp" not in _text(sheet.cell(bar_row, column).value).casefold():
                    continue
                for row in range(bar_row + 1, min(bar_row + 6, sheet.max_row + 1)):
                    name = _text(sheet.cell(row, column).value)
                    if name and name.casefold() not in {"scribed skills", "class masteries"}:
                        champion_points.append({"Name": name, "Points": ""})
            payload["ChampionPoints"] = champion_points

            settings_col = start_col + 4
            if settings_col <= sheet.max_column:
                for row in range(bar_row + 1, min(sheet.max_row, bar_row + 14) + 1):
                    setting = _text(sheet.cell(row, settings_col).value)
                    folded = setting.casefold()
                    if folded.startswith("attributes:"):
                        match = re.search(r"(\d+)\s*(health|mag(?:icka)?|stam(?:ina)?)", folded)
                        if match:
                            points = int(match.group(1))
                            resource = match.group(2)
                            if resource.startswith("health"):
                                payload["AttributeHealth"] = points
                            elif resource.startswith("mag"):
                                payload["AttributeMagicka"] = points
                            elif resource.startswith("stam"):
                                payload["AttributeStamina"] = points
                    elif folded.startswith("food:"):
                        payload["Food"] = setting.split(":", 1)[1].strip()
                    elif folded.startswith("mundus:"):
                        payload["Mundus"] = setting.split(":", 1)[1].strip()
                    elif folded.startswith("potion"):
                        payload["Potion"] = setting.split(":", 1)[1].strip() if ":" in setting else setting

        if not meaningful_gear and not any(payload["FrontBarSkills"]) and not any(payload["BackBarSkills"]):
            return None
        return ImportedBuildCandidate(
            gamertag=gamertag,
            build_name=build_name,
            eso_class=eso_class,
            role=role,
            assignment=role_text,
            source_sheet=sheet.title,
            payload=payload,
        )

    def _parse_sectioned_team_roster(self, sheet, plan: RosterImportPlan) -> None:
        section = ""
        for row in range(1, min(sheet.max_row, 150) + 1):
            cells = [_text(sheet.cell(row, col).value) for col in range(1, min(sheet.max_column, 8) + 1)]
            upper = {cell.upper() for cell in cells if cell}
            if "TANKS" in upper:
                section = "Tanks"
                continue
            if "HEALERS" in upper:
                section = "Healers"
                continue
            if "BUFF DPS" in upper:
                section = "Buff DPS"
                continue
            if any(label in upper for label in {"FULL DAMY DD", "FULL DAMAGE DD", "DPS"}):
                section = "Damage Dealers"
                continue
            if not section:
                continue

            player_cell = _text(sheet.cell(row, 2 if section == "Tanks" else 3).value)
            stack = "" if section == "Tanks" else _text(sheet.cell(row, 2).value)
            class_cell = _text(sheet.cell(row, 5).value)
            gear_cell = _text(sheet.cell(row, 6).value)
            if not player_cell or not class_cell:
                continue
            gamertag = player_cell.split(" - ", 1)[0].strip()
            if not gamertag or gamertag.casefold() in {"stack", "traditional and combat/crutch alerts numbers listed"}:
                continue
            eso_class = _class_from_text(class_cell)
            if not eso_class:
                continue
            role = "Tank" if section == "Tanks" else "Healer" if section == "Healers" else "Damage Dealer"
            assignment_parts = []
            if " - " in player_cell:
                assignment_parts.append(player_cell.split(" - ", 1)[1].strip())
            if stack and stack.casefold() != "stack":
                assignment_parts.append(f"Stack {stack}")
            if section == "Buff DPS":
                assignment_parts.append("Buff DPS")
            assignment = " · ".join(part for part in assignment_parts if part)
            member = self._ensure_member(
                plan,
                gamertag=gamertag,
                eso_class=eso_class,
                role=role,
                assignment=assignment,
                source=f"{sheet.title}: {section}",
                authoritative_role=True,
            )
            if gear_cell:
                member.warnings.append(f"Generalized gear note available: {gear_cell}")

    def _attach_role_templates(self, workbook, plan: RosterImportPlan) -> None:
        for sheet in workbook.worksheets:
            if sheet.title.casefold() == "team roster":
                continue
            template = self._parse_role_template(sheet)
            if template is None:
                continue
            matches = [member for member in plan.members if self._template_matches_member(sheet.title, member)]
            if not matches:
                plan.warnings.append(f"No roster player matched build template {sheet.title}.")
                continue
            for member in matches:
                candidate = deepcopy(template)
                candidate.gamertag = member.gamertag
                candidate.eso_class = member.eso_class or candidate.eso_class
                candidate.role = member.primary_role or candidate.role
                candidate.assignment = member.assignment
                candidate.payload["Gamertag"] = member.gamertag
                candidate.payload["EsoClass"] = candidate.eso_class
                candidate.payload["Role"] = candidate.role
                member.builds.append(candidate)

    def _parse_role_template(self, sheet) -> ImportedBuildCandidate | None:
        title = _text(sheet.title)
        eso_class = _class_from_text(title)
        role = _role_from_text(title, default_damage=True)
        if title.casefold() in {"main tank", "off tank"}:
            role = "Tank"
        if "healer" in title.casefold():
            role = "Healer"
        if not eso_class:
            for row in range(1, min(sheet.max_row, 20) + 1):
                if _text(sheet.cell(row, 2).value).casefold() == "pure class":
                    values = " ".join(
                        _text(sheet.cell(row, col).value)
                        for col in range(4, min(sheet.max_column, 8) + 1)
                    )
                    eso_class = _class_from_text(values)
                    break
        if not eso_class and role != "Tank":
            return None

        fields: dict[str, str] = {}
        for row in range(1, min(sheet.max_row, 120) + 1):
            label = _text(sheet.cell(row, 2).value)
            if not label:
                continue
            value = " | ".join(
                piece
                for piece in (
                    _text(sheet.cell(row, col).value)
                    for col in range(4, min(sheet.max_column, 9) + 1)
                )
                if piece
            )
            if value:
                fields[label.casefold()] = value

        front = next((value for key, value in fields.items() if "front bar skills" in key), "")
        back = next((value for key, value in fields.items() if "back bar skills" in key), "")
        gear = fields.get("gear", "")
        if not front and not back and not gear:
            return None

        champion_points: list[dict[str, str]] = []
        for key, value in fields.items():
            if key in {"blue cp", "red cp"}:
                for entry in re.split(r"\s*\|\s*|\s*,\s*", value):
                    entry = entry.strip()
                    if entry:
                        champion_points.append({"Name": entry, "Points": ""})

        notes = []
        if gear:
            notes.append(f"Gear summary: {gear}")
        if fields.get("pure class"):
            notes.append(f"Pure class/passives: {fields['pure class']}")
        for key in ("additional info", "flex gear setup", "flex skills", "navi info"):
            if fields.get(key):
                notes.append(f"{key.title()}: {fields[key]}")
        return ImportedBuildCandidate(
            gamertag="",
            build_name=title,
            eso_class=eso_class,
            role=role,
            source_sheet=sheet.title,
            payload={
                "BuildName": title,
                "EsoClass": eso_class,
                "Role": role,
                "FrontBarSkills": _split_skills(front),
                "BackBarSkills": _split_skills(back),
                "ChampionPoints": champion_points,
                "Notes": "\n".join(notes),
                "Werewolf": "ww" in title.casefold() or "werewolf" in title.casefold(),
            },
        )

    @staticmethod
    def _template_matches_member(template_title: str, member: ImportedRosterMember) -> bool:
        title = template_title.casefold()
        role = member.primary_role.casefold()
        eso_class = member.eso_class.casefold()
        assignment = member.assignment.casefold()
        source = member.source.casefold()
        if title == "main tank":
            return role == "tank" and ("main" in assignment or "main" in source)
        if title == "off tank":
            return role == "tank" and not ("main" in assignment or "main" in source)
        if "healer" in title and role != "healer":
            return False
        if any(token in title for token in (" dd", "zens", "kosh", "ww")) and role != "damage dealer":
            return False
        if "zens" in title or "kosh" in title:
            return eso_class == "dragonknight" or "zen" in assignment
        if "ww" in title:
            return eso_class == "nightblade" and ("ww" in assignment or "ww" in source or "werewolf" in assignment)
        template_class = _class_from_text(template_title)
        return bool(template_class and template_class.casefold() == eso_class)


def _existing_character_map(roster_service, build_service: BuildService) -> dict[str, list[tuple[str, str]]]:
    result: dict[str, list[tuple[str, str]]] = {}
    catalog = build_service.canonical.catalog_service.load()
    players = {
        _text(player.get("player_id")): _text(player.get("gamertag"))
        for player in catalog.get("players", [])
        if isinstance(player, dict)
    }
    for character in catalog.get("characters", []):
        if not isinstance(character, dict):
            continue
        gamertag = players.get(_text(character.get("player_id")), _text(character.get("gamertag")))
        name = _text(character.get("name"))
        eso_class = _normalize_class(character.get("eso_class"))
        if gamertag and name:
            result.setdefault(gamertag.casefold(), []).append((name, eso_class))
    for member in roster_service.list_members():
        gamertag = _text(member.PlayerName)
        name = _text(member.CharacterName)
        eso_class = _normalize_class(member.EsoClass)
        if gamertag and name:
            pair = (name, eso_class)
            bucket = result.setdefault(gamertag.casefold(), [])
            if pair not in bucket:
                bucket.append(pair)
    return result


def resolve_import_characters(plan: RosterImportPlan, roster_service, build_service: BuildService) -> None:
    known = _existing_character_map(roster_service, build_service)
    for member in plan.members:
        if member.character_name:
            continue
        candidates = known.get(member.gamertag.casefold(), [])
        same_class = [
            name
            for name, eso_class in candidates
            if member.eso_class and eso_class.casefold() == member.eso_class.casefold()
        ]
        if len(same_class) == 1:
            member.character_name = same_class[0]
            continue
        unique_names = sorted({name for name, _ in candidates}, key=str.casefold)
        if len(unique_names) == 1:
            member.character_name = unique_names[0]
        elif member.builds:
            member.warnings.append(
                "Build detected, but character identity needs confirmation before it can be imported."
            )


def _merge_team_names(existing: str, team_name: str) -> str:
    names: list[str] = []
    seen: set[str] = set()
    for raw in f"{existing},{team_name}".split(","):
        name = raw.strip()
        key = name.casefold()
        if name and key not in seen:
            names.append(name)
            seen.add(key)
    return ", ".join(names)


def apply_roster_import(
    plan: RosterImportPlan,
    roster_service,
    build_service: BuildService,
    *,
    import_builds: bool = True,
) -> RosterImportResult:
    """Merge reviewed rows; never delete unrelated roster members or builds."""

    team_name = _text(plan.team_name)
    if not team_name:
        raise ValueError("Team name is required before importing a roster.")
    canonical_team = roster_service.ensure_team_name(team_name)
    existing_members = roster_service.list_members()
    selected_members = [member for member in plan.members if member.selected]
    created = 0
    updated = 0
    warnings: list[str] = []

    for imported in selected_members:
        tag_key = imported.gamertag.casefold()
        character_key = imported.character_name.casefold()
        matches = [
            member
            for member in existing_members
            if member.PlayerName.casefold() == tag_key
            and (
                not character_key
                or not member.CharacterName
                or member.CharacterName.casefold() == character_key
            )
        ]
        target = matches[0] if len(matches) == 1 else None
        if target is None and not character_key:
            same_player = [member for member in existing_members if member.PlayerName.casefold() == tag_key]
            target = same_player[0] if len(same_player) == 1 else None
        if target is None:
            target = RosterMember(
                PlayerName=imported.gamertag,
                CharacterName=imported.character_name,
                EsoClass=imported.eso_class,
                PrimaryRole=imported.primary_role,
                SecondaryRole=imported.secondary_role,
                Status=imported.status or "Active",
                Team=canonical_team,
            )
            target.Id = roster_service.create_member(target)
            existing_members.append(target)
            created += 1
        else:
            target.CharacterName = imported.character_name or target.CharacterName
            target.EsoClass = imported.eso_class or target.EsoClass
            target.PrimaryRole = imported.primary_role or target.PrimaryRole
            target.SecondaryRole = imported.secondary_role or target.SecondaryRole
            target.Status = imported.status or target.Status or "Active"
            target.Team = _merge_team_names(target.Team, canonical_team)
            roster_service.update_member(target)
            updated += 1

    imported_build_count = 0
    skipped_builds = 0
    if import_builds:
        roster = build_service.load()
        saved_builds = [
            member
            for member in roster.Members
            if member.Name or member.Gamertag or member.BuildName
        ]
        pending_assignments: list[tuple[str, str, str, str, str]] = []
        for imported in selected_members:
            for candidate in imported.builds:
                if not imported.character_name:
                    skipped_builds += 1
                    warnings.append(
                        f"Skipped {candidate.build_name} for {imported.gamertag}: choose a character in the import preview first."
                    )
                    continue
                payload = deepcopy(candidate.payload)
                payload["Name"] = imported.character_name
                payload["Gamertag"] = imported.gamertag
                payload["BuildName"] = candidate.build_name
                payload["EsoClass"] = imported.eso_class or candidate.eso_class
                payload["Role"] = imported.primary_role or candidate.role
                build = PlayerBuild.from_dict(payload)
                key = (
                    build.Gamertag.casefold(),
                    build.Name.casefold(),
                    build.BuildName.casefold(),
                )
                replacement = next(
                    (
                        index
                        for index, existing in enumerate(saved_builds)
                        if (
                            existing.Gamertag.casefold(),
                            existing.Name.casefold(),
                            existing.BuildName.casefold(),
                        ) == key
                    ),
                    None,
                )
                if replacement is None:
                    saved_builds.append(build)
                else:
                    saved_builds[replacement] = build
                pending_assignments.append(
                    (
                        build.Gamertag,
                        build.Name,
                        build.BuildName,
                        imported.primary_role,
                        candidate.assignment or imported.assignment,
                    )
                )
                imported_build_count += 1

        if imported_build_count:
            build_service.save(BuildRoster(Members=saved_builds))
            catalog_service = build_service.canonical.catalog_service
            catalog = catalog_service.load()
            players_by_id = {
                _text(player.get("player_id")): player
                for player in catalog.get("players", [])
                if isinstance(player, dict)
            }
            for gamertag, character_name, build_name, raid_role, assignment in pending_assignments:
                player_ids = {
                    player_id
                    for player_id, player in players_by_id.items()
                    if _text(player.get("gamertag")).casefold() == gamertag.casefold()
                }
                characters = [
                    character
                    for character in catalog.get("characters", [])
                    if isinstance(character, dict)
                    and _text(character.get("player_id")) in player_ids
                    and _text(character.get("name")).casefold() == character_name.casefold()
                ]
                if len(characters) != 1:
                    warnings.append(
                        f"Imported {build_name}, but could not uniquely attach its team assignment for {gamertag}."
                    )
                    continue
                character_id = _text(characters[0].get("character_id"))
                builds = [
                    build
                    for build in catalog.get("builds", [])
                    if isinstance(build, dict)
                    and _text(build.get("character_id")) == character_id
                    and _text(build.get("name")).casefold() == build_name.casefold()
                ]
                if len(builds) != 1:
                    warnings.append(
                        f"Imported {build_name}, but its canonical build identity could not be reloaded uniquely."
                    )
                    continue
                catalog_service.assign_build_to_team(
                    build_id=_text(builds[0].get("build_id")),
                    team_name=canonical_team,
                    raid_role=raid_role,
                    slot_name=assignment,
                    notes="Imported from roster workbook",
                )

    return RosterImportResult(
        created_roster_members=created,
        updated_roster_members=updated,
        imported_builds=imported_build_count,
        skipped_builds=skipped_builds,
        warnings=tuple(warnings),
    )


class RosterImportPreviewDialog(QDialog):
    def __init__(self, plan: RosterImportPlan, parent=None):
        super().__init__(parent)
        self.plan = plan
        self.setWindowTitle("Import Roster")
        self.resize(1050, 650)
        root = QVBoxLayout(self)

        title = QLabel(f"Preview: {plan.source_path.name}")
        title.setProperty("pageTitle", True)
        root.addWidget(title)
        summary = QLabel(
            f"Detected {len(plan.members)} player(s) and {plan.build_count} build candidate(s). "
            "Nothing is written until you confirm this preview."
        )
        summary.setWordWrap(True)
        root.addWidget(summary)

        team_row = QHBoxLayout()
        team_row.addWidget(QLabel("TEAM"))
        self.team_edit = QLineEdit(plan.team_name)
        team_row.addWidget(self.team_edit, 1)
        self.import_builds = QCheckBox("Import detected builds")
        self.import_builds.setChecked(True)
        team_row.addWidget(self.import_builds)
        root.addLayout(team_row)

        self.table = QTableWidget(len(plan.members), 8)
        self.table.setHorizontalHeaderLabels(
            ["Import", "Gamertag", "Character", "Class", "Role", "Assignment", "Builds", "Source"]
        )
        self.table.horizontalHeader().setStretchLastSection(True)
        for row, member in enumerate(plan.members):
            enabled = QTableWidgetItem("")
            enabled.setFlags(enabled.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            enabled.setCheckState(Qt.CheckState.Checked if member.selected else Qt.CheckState.Unchecked)
            self.table.setItem(row, 0, enabled)
            values = (
                member.gamertag,
                member.character_name,
                member.eso_class,
                member.primary_role,
                member.assignment,
                str(len(member.builds)),
                member.source,
            )
            for column, value in enumerate(values, start=1):
                item = QTableWidgetItem(value)
                if column != 2:
                    item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row, column, item)
        self.table.resizeColumnsToContents()
        root.addWidget(self.table, 1)

        warning_lines = list(plan.warnings)
        warning_lines.extend(
            f"{member.gamertag}: {warning}"
            for member in plan.members
            for warning in member.warnings
        )
        warnings = QLabel("\n".join(warning_lines) if warning_lines else "No parser warnings.")
        warnings.setWordWrap(True)
        warnings.setProperty("muted", True)
        root.addWidget(warnings)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Cancel | QDialogButtonBox.StandardButton.Ok
        )
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("Import Selected")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

    def sync_to_plan(self) -> None:
        self.plan.team_name = self.team_edit.text().strip()
        for row, member in enumerate(self.plan.members):
            enabled = self.table.item(row, 0)
            member.selected = bool(enabled and enabled.checkState() == Qt.CheckState.Checked)
            character = self.table.item(row, 2)
            member.character_name = character.text().strip() if character is not None else ""


def _import_roster_from_file(page) -> None:
    filename, _ = QFileDialog.getOpenFileName(
        page,
        "Import Roster",
        "",
        "Roster files (*.xlsx *.xlsm *.csv *.json);;Excel workbooks (*.xlsx *.xlsm);;CSV (*.csv);;JSON (*.json)",
    )
    if not filename:
        return
    try:
        plan = RosterImportParser().parse_file(Path(filename))
        build_service = BuildService(get_data_dir() / "builds.json")
        build_service.load()
        resolve_import_characters(plan, page.roster_service, build_service)
        dialog = RosterImportPreviewDialog(plan, page)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        dialog.sync_to_plan()
        if not [member for member in plan.members if member.selected]:
            page.status.warning("No roster rows were selected for import.")
            return
        if not plan.team_name:
            page.status.warning("Enter a team name before importing.")
            return
        UserSafetySnapshotService().create(
            f"roster-import-{Path(filename).stem}"
        )
        result = apply_roster_import(
            plan,
            page.roster_service,
            build_service,
            import_builds=dialog.import_builds.isChecked(),
        )
        page.refresh()
        message = (
            f"Roster import complete: {result.created_roster_members} new, "
            f"{result.updated_roster_members} updated, {result.imported_builds} build(s) imported."
        )
        if result.skipped_builds:
            message += (
                f" {result.skipped_builds} build(s) were skipped because character identity was unresolved."
            )
        page.status.success(message)
        if result.warnings:
            QMessageBox.information(page, "Roster Import Notes", "\n".join(result.warnings))
    except Exception as exc:
        page.status.error(f"Roster import failed: {exc}")
        QMessageBox.critical(page, "Roster Import Failed", str(exc))


def install() -> None:
    global _INSTALLED
    if _INSTALLED:
        return
    from ui.themed_roster_page import RosterPage

    original_build_assignments_tab = RosterPage._build_assignments_tab

    def build_assignments_tab_with_import(self):
        page = original_build_assignments_tab(self)
        for button in page.findChildren(QPushButton):
            if button.text().strip().casefold() == "import roster":
                self.import_roster_button = button
                button.setToolTip(
                    "Import a raid roster from Excel, CSV, or JSON. Preview players and detected builds before saving."
                )
                button.clicked.connect(lambda _checked=False: _import_roster_from_file(self))
                break
        return page

    RosterPage._build_assignments_tab = build_assignments_tab_with_import
    _INSTALLED = True
