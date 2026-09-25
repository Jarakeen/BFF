from __future__ import annotations

"""Raid Plan scoped Build export helpers.

Raid Plans own the ordered chair selection; canonical BuildService owns build payloads.
This module joins them for sharing without inventing another persistence store.
"""

import csv
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

from models.build_model import BuildRoster, PlayerBuild
from models.raid_plan import RaidPlan, RaidPlanMember
from services.build_service import BuildService


def _clean(value: object) -> str:
    return " ".join(str(value or "").strip().split())


@dataclass(frozen=True)
class RaidPlanBuildExport:
    roster: BuildRoster
    seats: tuple[tuple[RaidPlanMember, PlayerBuild], ...]
    unresolved_seats: tuple[str, ...] = ()


def _fallback_matches(
    member: RaidPlanMember,
    builds: tuple[PlayerBuild, ...],
) -> tuple[PlayerBuild, ...]:
    wanted_player = _clean(member.gamertag).casefold()
    wanted_character = _clean(member.character_name).casefold()
    wanted_name = _clean(member.selected_build_name).casefold()

    matches = []
    for build in builds:
        if wanted_player and _clean(build.Gamertag).casefold() != wanted_player:
            continue
        if wanted_character and _clean(build.Name).casefold() != wanted_character:
            continue
        if wanted_name and _clean(build.BuildName).casefold() != wanted_name:
            continue
        matches.append(build)
    return tuple(matches)


def raid_plan_build_export(
    plan: RaidPlan,
    build_service: BuildService,
) -> RaidPlanBuildExport:
    """Resolve exact canonical builds for one Raid Plan, preserving seat order."""

    builds = tuple(build_service.load().Members)
    by_id = {
        _clean(getattr(build, "BuildId", "")): build
        for build in builds
        if _clean(getattr(build, "BuildId", ""))
    }

    resolved: list[tuple[RaidPlanMember, PlayerBuild]] = []
    unresolved: list[str] = []
    seen_ids: set[str] = set()

    for member in plan.members:
        selected_id = _clean(member.selected_build_id)
        build = by_id.get(selected_id) if selected_id else None

        if build is None and not selected_id and member.selected_build_name:
            matches = _fallback_matches(member, builds)
            build = matches[0] if len(matches) == 1 else None

        if build is None:
            if member.build_selected or member.planned_gear_sets or member.planned_skills:
                unresolved.append(member.seat_id)
            continue

        build_id = _clean(getattr(build, "BuildId", ""))
        if build_id and build_id in seen_ids:
            continue
        if build_id:
            seen_ids.add(build_id)
        resolved.append((member, build))

    return RaidPlanBuildExport(
        roster=BuildRoster(Members=[build for _, build in resolved]),
        seats=tuple(resolved),
        unresolved_seats=tuple(unresolved),
    )


def _planned_or_saved_sets(member: RaidPlanMember, build: PlayerBuild) -> tuple[str, ...]:
    planned = tuple(_clean(value) for value in member.planned_gear_sets if _clean(value))
    if planned:
        return planned

    values: list[str] = []
    for slot in build.Armor.values():
        if isinstance(slot, dict):
            values.extend((_clean(slot.get("Set")), _clean(slot.get("Set2"))))
    for slot in (
        build.FrontBarWeapon,
        build.FrontBarOffHand,
        build.BackBarWeapon,
        build.BackBarOffHand,
        build.Necklace,
        build.Ring1,
        build.Ring2,
    ):
        values.extend((_clean(getattr(slot, "Set", "")), _clean(getattr(slot, "Set2", ""))))
    return tuple(dict.fromkeys(value for value in values if value))


def raid_plan_discord_builds_text(
    plan: RaidPlan,
    export: RaidPlanBuildExport,
) -> str:
    """Format the Raid Plan's exact resolved builds for Discord clipboard sharing."""

    difficulty = _clean(plan.difficulty) or "Difficulty not set"
    lines = [
        f"**{_clean(plan.name)}**",
        f"{_clean(plan.trial_id)} • {difficulty}",
        "",
        "**Builds**",
    ]

    for member, build in export.seats:
        player = _clean(member.gamertag) or _clean(build.Gamertag) or "Open"
        character = _clean(member.character_name) or _clean(build.Name)
        eso_class = _clean(member.eso_class) or _clean(build.EsoClass)
        build_name = _clean(build.BuildName) or _clean(member.selected_build_name) or "Build"
        identity = " • ".join(value for value in (player, character, eso_class) if value)
        lines.append(f"**{member.seat_id}** | {identity}")
        lines.append(f"↳ {build_name}")

        sets = _planned_or_saved_sets(member, build)
        if sets:
            lines.append("↳ Sets: " + " + ".join(sets))
        skills = tuple(_clean(value) for value in member.planned_skills if _clean(value))
        if skills:
            lines.append("↳ Planned skills: " + ", ".join(skills))

    if export.unresolved_seats:
        lines.extend(
            (
                "",
                "⚠ Unresolved build seats: " + ", ".join(export.unresolved_seats),
            )
        )

    return "\n".join(lines).strip()


def export_raid_plan_builds_csv(
    plan: RaidPlan,
    export: RaidPlanBuildExport,
    path: str | Path,
) -> Path:
    """Write a compact, spreadsheet-friendly Raid Plan build sheet."""

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            (
                "Seat",
                "Player",
                "Character",
                "Class",
                "Role",
                "Build",
                "Sets",
                "Planned Skills",
                "Mundus",
                "Primary Assignment",
                "Utility Assignments",
            )
        )
        for member, build in export.seats:
            writer.writerow(
                (
                    member.seat_id,
                    _clean(member.gamertag) or _clean(build.Gamertag),
                    _clean(member.character_name) or _clean(build.Name),
                    _clean(member.eso_class) or _clean(build.EsoClass),
                    _clean(member.role) or _clean(build.Role),
                    _clean(build.BuildName) or _clean(member.selected_build_name),
                    " + ".join(_planned_or_saved_sets(member, build)),
                    ", ".join(_clean(value) for value in member.planned_skills if _clean(value)),
                    _clean(member.planned_mundus) or _clean(build.Mundus),
                    _clean(member.primary_assignment),
                    ", ".join(_clean(value) for value in member.utility_assignments if _clean(value)),
                )
            )
    return target


def _gear_rows(build: PlayerBuild) -> tuple[tuple[str, str, str, str, str], ...]:
    """Return slot, item/set, type/weight, trait, enchant without baseline level/quality noise."""
    rows: list[tuple[str, str, str, str, str]] = []
    for slot_name, values in build.Armor.items():
        values = values if isinstance(values, dict) else {}
        set_name = " + ".join(
            value for value in (_clean(values.get("Set")), _clean(values.get("Set2"))) if value
        )
        rows.append((
            slot_name,
            set_name or "—",
            _clean(values.get("Weight")) or "—",
            _clean(values.get("Trait")) or "—",
            _clean(values.get("Enchant")) or "—",
        ))
    for slot_name, gear in (
        ("Front Bar", build.FrontBarWeapon),
        ("Front Off Hand", build.FrontBarOffHand),
        ("Back Bar", build.BackBarWeapon),
        ("Back Off Hand", build.BackBarOffHand),
        ("Necklace", build.Necklace),
        ("Ring 1", build.Ring1),
        ("Ring 2", build.Ring2),
    ):
        if getattr(gear, "is_empty", False) and slot_name in {"Front Off Hand", "Back Off Hand"}:
            continue
        set_name = " + ".join(
            value for value in (_clean(gear.Set), _clean(gear.Set2)) if value
        )
        kind = _clean(gear.WeaponType) or _clean(gear.Weight)
        rows.append((
            slot_name,
            set_name or "—",
            kind or "—",
            _clean(gear.Trait) or "—",
            _clean(gear.Enchant) or "—",
        ))
    return tuple(rows)


def _skill_rows(build: PlayerBuild) -> tuple[tuple[str, str, str], ...]:
    rows: list[tuple[str, str, str]] = []
    for bar_name, skills in (("Front", build.FrontBarSkills), ("Back", build.BackBarSkills)):
        for index, skill in enumerate(skills):
            name = _clean(skill)
            if name:
                position = "Ultimate" if index >= 5 else str(index + 1)
                rows.append((bar_name, position, name))
    return tuple(rows)


def _scribed_rows(build: PlayerBuild) -> tuple[tuple[str, str, str, str, str], ...]:
    recipes = tuple(
        recipe for recipe in build.ScribedSkillRecipes if _clean(getattr(recipe, "ResultName", ""))
    )
    if recipes:
        return tuple(
            (
                _clean(recipe.ResultName),
                _clean(recipe.Grimoire) or "—",
                _clean(recipe.Focus) or "—",
                _clean(recipe.Signature) or "—",
                _clean(recipe.Affix) or "—",
            )
            for recipe in recipes
        )
    return tuple((_clean(name), "—", "—", "—", "—") for name in build.ScribedSkills if _clean(name))


def export_raid_plan_builds_xlsx(
    plan: RaidPlan,
    export: RaidPlanBuildExport,
    path: str | Path,
) -> Path:
    """Write a low-ink FoundryDock workbook with a raid index and one detailed sheet per player."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("XLSX export requires openpyxl.") from exc

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    index = wb.active
    index.title = "Raid Index"
    teal, dark, gold, pale, gray = "1F3F45", "303A3C", "C8A46A", "E5ECEB", "687476"
    hair = Side(style="hair", color="C9D1D0")
    rule = Side(style="thin", color=gold)

    def safe(value: object) -> str:
        return _clean(value)

    def title_block(ws, title: str, subtitle: str, last_col: int = 6) -> None:
        last = get_column_letter(last_col)
        ws.merge_cells(f"A1:{last}1")
        ws["A1"] = title
        ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=teal)
        ws.merge_cells(f"A2:{last}2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color=gray)
        ws["A3"].border = Border(bottom=rule)
        for col in range(2, last_col + 1):
            ws.cell(3, col).border = Border(bottom=rule)
        ws.row_dimensions[1].height = 27
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A4"
        ws.oddFooter.center.text = "FoundryDock • Leave Better Records"

    def section(ws, row: int, label: str, last_col: int = 6) -> int:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        cell = ws.cell(row, 1, label.upper())
        cell.font = Font(name="Calibri", size=10, bold=True, color=teal)
        cell.fill = PatternFill("solid", fgColor=pale)
        cell.border = Border(bottom=rule)
        return row + 1

    title_block(
        index,
        safe(plan.name) or "Raid Plan",
        f"{safe(plan.trial_id)}  •  {safe(plan.difficulty) or 'Difficulty not set'}  •  {len(export.seats)} linked Builds",
        7,
    )
    headers = ("Seat", "Player", "Character", "Class", "Role", "Build", "Sheet")
    for col, heading in enumerate(headers, 1):
        c = index.cell(5, col, heading)
        c.font = Font(bold=True, color=teal)
        c.border = Border(bottom=rule)
    used_names: set[str] = {"Raid Index", "Unresolved"}
    player_sheets: list[tuple[RaidPlanMember, PlayerBuild, str]] = []
    for row_no, (member, build) in enumerate(export.seats, 6):
        base = safe(member.gamertag) or safe(build.Gamertag) or safe(member.seat_id) or "Build"
        base = "".join(ch for ch in base if ch not in r'[]:*?/\\')[:28] or "Build"
        sheet_name = base
        suffix = 2
        while sheet_name.casefold() in {name.casefold() for name in used_names}:
            sheet_name = f"{base[:25]} {suffix}"
            suffix += 1
        used_names.add(sheet_name)
        player_sheets.append((member, build, sheet_name))
        vals = (
            member.seat_id, member.gamertag or build.Gamertag, member.character_name or build.Name,
            member.eso_class or build.EsoClass, member.role or build.Role,
            build.BuildName or member.selected_build_name, sheet_name,
        )
        for col, value in enumerate(vals, 1):
            c = index.cell(row_no, col, safe(value))
            c.data_type = "s"
            c.font = Font(size=10, color=dark)
            c.border = Border(bottom=hair)
        index.cell(row_no, 7).hyperlink = f"#'{sheet_name}'!A1"
        index.cell(row_no, 7).style = "Hyperlink"
    for col, width in enumerate((13, 20, 22, 18, 14, 31, 22), 1):
        index.column_dimensions[get_column_letter(col)].width = width
    index.auto_filter.ref = f"A5:G{max(5, index.max_row)}"
    index.freeze_panes = "A6"

    for member, build, sheet_name in player_sheets:
        ws = wb.create_sheet(sheet_name)
        identity = " • ".join(
            value for value in (
                safe(member.seat_id),
                safe(member.gamertag) or safe(build.Gamertag),
                safe(member.character_name) or safe(build.Name),
                safe(member.eso_class) or safe(build.EsoClass),
            ) if value
        )
        title_block(ws, safe(build.BuildName) or safe(member.selected_build_name) or "Build", identity, 6)
        row = 5
        row = section(ws, row, "Build Summary")
        summary = (
            ("Role", safe(member.role) or safe(build.Role) or "—", "Mundus", safe(member.planned_mundus) or safe(build.Mundus) or "—"),
            ("Food", safe(build.Food) or "—", "Potion", safe(build.Potion) or "—"),
            ("Primary Assignment", safe(member.primary_assignment) or "—", "Utility", ", ".join(member.utility_assignments) or "—"),
        )
        for left_label, left_value, right_label, right_value in summary:
            ws.cell(row, 1, left_label).font = Font(bold=True, color=teal)
            ws.cell(row, 2, left_value).data_type = "s"
            ws.cell(row, 4, right_label).font = Font(bold=True, color=teal)
            ws.cell(row, 5, right_value).data_type = "s"
            row += 1
        row += 1

        row = section(ws, row, "Gear by Slot")
        gear_headers = ("Slot", "Set / Item", "Weight / Type", "Trait", "Enchant")
        for col, heading in enumerate(gear_headers, 1):
            ws.cell(row, col, heading).font = Font(bold=True, color=teal)
            ws.cell(row, col).border = Border(bottom=rule)
        row += 1
        for gear_row in _gear_rows(build):
            for col, value in enumerate(gear_row, 1):
                c = ws.cell(row, col, safe(value))
                c.data_type = "s"
                c.border = Border(bottom=hair)
                c.alignment = Alignment(vertical="top", wrap_text=True)
            row += 1
        row += 1

        row = section(ws, row, "Skill Bars")
        for col, heading in enumerate(("Bar", "Slot", "Skill"), 1):
            ws.cell(row, col, heading).font = Font(bold=True, color=teal)
            ws.cell(row, col).border = Border(bottom=rule)
        row += 1
        for skill_row in _skill_rows(build):
            for col, value in enumerate(skill_row, 1):
                c = ws.cell(row, col, safe(value))
                c.data_type = "s"
                c.border = Border(bottom=hair)
            row += 1
        row += 1

        row = section(ws, row, "Scribed Skill Formulas")
        for col, heading in enumerate(("Result", "Grimoire", "Focus", "Signature", "Affix"), 1):
            ws.cell(row, col, heading).font = Font(bold=True, color=teal)
            ws.cell(row, col).border = Border(bottom=rule)
        row += 1
        scribed = _scribed_rows(build)
        if not scribed:
            ws.cell(row, 1, "No scribed skills recorded.")
            row += 1
        else:
            for recipe in scribed:
                for col, value in enumerate(recipe, 1):
                    ws.cell(row, col, safe(value)).border = Border(bottom=hair)
                row += 1

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 34
        ws.column_dimensions["F"].width = 4
        for cells in ws.iter_rows():
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True

    if export.unresolved_seats:
        gaps = wb.create_sheet("Unresolved")
        title_block(gaps, "Unresolved Build Seats", "These chairs do not have an exact linked Saved Build.", 2)
        gaps.append(())
        gaps.append(("Seat", "Next step"))
        for seat in export.unresolved_seats:
            gaps.append((seat, "Review the chair in Roles and link the intended Saved Build."))
        gaps.column_dimensions["A"].width = 20
        gaps.column_dimensions["B"].width = 70

    wb.save(target)
    wb.close()
    return target


def export_raid_plan_builds_pdf(
    plan: RaidPlan,
    export: RaidPlanBuildExport,
    path: str | Path,
) -> Path:
    """Write a thematic, low-ink build packet with one player section per page."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("PDF export requires ReportLab. The packaged FoundryDock build includes it.") from exc

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    teal = colors.HexColor("#1F3F45")
    gold = colors.HexColor("#C8A46A")
    dark = colors.HexColor("#303A3C")
    gray = colors.HexColor("#687476")
    pale = colors.HexColor("#E5ECEB")
    hair = colors.HexColor("#D9DFDE")

    title = ParagraphStyle("FDTitle", fontName="Helvetica-Bold", fontSize=18, leading=21, textColor=teal, spaceAfter=2)
    subtitle = ParagraphStyle("FDSub", fontName="Helvetica", fontSize=8.5, leading=11, textColor=gray, spaceAfter=8)
    section_style = ParagraphStyle("FDSection", fontName="Helvetica-Bold", fontSize=9, leading=11, textColor=teal, spaceBefore=5, spaceAfter=4)
    body = ParagraphStyle("FDBody", fontName="Helvetica", fontSize=7.7, leading=9.5, textColor=dark)
    small = ParagraphStyle("FDSmall", fontName="Helvetica", fontSize=6.8, leading=8.2, textColor=gray)

    def P(value: object, style=body) -> Paragraph:
        return Paragraph(escape(_clean(value) or "—"), style)

    def table(rows, widths, *, header=True):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
        commands = [
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LINEBELOW", (0, 0), (-1, -1), 0.2, hair),
        ]
        if header:
            commands.extend([
                ("BACKGROUND", (0, 0), (-1, 0), pale),
                ("LINEBELOW", (0, 0), (-1, 0), 0.7, gold),
            ])
        t.setStyle(TableStyle(commands))
        return t

    doc = SimpleDocTemplate(
        str(target), pagesize=LETTER, leftMargin=.48*inch, rightMargin=.48*inch,
        topMargin=.45*inch, bottomMargin=.45*inch,
        title=f"{_clean(plan.name)} Raid Plan Builds",
        author="FoundryDock",
    )
    story = [
        Paragraph(escape(_clean(plan.name) or "Raid Plan"), title),
        Paragraph(
            f"{escape(_clean(plan.trial_id))} · {escape(_clean(plan.difficulty) or 'Difficulty not set')} · "
            f"{len(export.seats)} linked Builds<br/><i>FoundryDock field packet · Leave Better Records</i>",
            subtitle,
        ),
    ]
    index_rows = [[P("Seat", section_style), P("Player", section_style), P("Character", section_style), P("Build", section_style)]]
    for member, build in export.seats:
        index_rows.append([
            P(member.seat_id), P(member.gamertag or build.Gamertag),
            P(member.character_name or build.Name), P(build.BuildName or member.selected_build_name),
        ])
    story.extend([table(index_rows, (.75*inch, 1.55*inch, 1.75*inch, 2.5*inch)), Spacer(1, 8)])
    if export.unresolved_seats:
        story.append(Paragraph("Unresolved chairs: " + escape(", ".join(export.unresolved_seats)), subtitle))

    for member, build in export.seats:
        story.append(PageBreak())
        identity = " • ".join(value for value in (
            _clean(member.seat_id), _clean(member.gamertag) or _clean(build.Gamertag),
            _clean(member.character_name) or _clean(build.Name),
            _clean(member.eso_class) or _clean(build.EsoClass),
        ) if value)
        story.append(Paragraph(escape(_clean(build.BuildName) or _clean(member.selected_build_name) or "Build"), title))
        story.append(Paragraph(escape(identity), subtitle))

        summary_rows = [
            [P("Role", section_style), P(member.role or build.Role), P("Mundus", section_style), P(member.planned_mundus or build.Mundus)],
            [P("Food", section_style), P(build.Food), P("Potion", section_style), P(build.Potion)],
            [P("Assignment", section_style), P(member.primary_assignment), P("Utility", section_style), P(", ".join(member.utility_assignments))],
        ]
        story.extend([table(summary_rows, (.75*inch, 2.6*inch, .75*inch, 2.6*inch), header=False), Spacer(1, 5)])

        story.append(Paragraph("GEAR BY SLOT", section_style))
        gear_rows = [[P(x, section_style) for x in ("Slot", "Set / Item", "Weight / Type", "Trait", "Enchant")]]
        gear_rows.extend([[P(value) for value in row] for row in _gear_rows(build)])
        story.append(table(gear_rows, (.85*inch, 2.15*inch, 1.05*inch, 1.15*inch, 1.55*inch)))

        story.append(Paragraph("SKILL BARS", section_style))
        skill_rows = [[P(x, section_style) for x in ("Bar", "Slot", "Skill")]]
        skill_rows.extend([[P(value) for value in row] for row in _skill_rows(build)])
        story.append(table(skill_rows, (.75*inch, .7*inch, 5.3*inch)))

        story.append(Paragraph("SCRIBED SKILL FORMULAS", section_style))
        recipes = _scribed_rows(build)
        if recipes:
            recipe_rows = [[P(x, section_style) for x in ("Result", "Grimoire", "Focus", "Signature", "Affix")]]
            recipe_rows.extend([[P(value) for value in row] for row in recipes])
            story.append(table(recipe_rows, (1.35*inch, 1.25*inch, 1.35*inch, 1.35*inch, 1.45*inch)))
        else:
            story.append(Paragraph("No scribed skills recorded.", small))

    doc.build(story)
    return target


__all__ = [
    "RaidPlanBuildExport",
    "export_raid_plan_builds_csv",
    "export_raid_plan_builds_pdf",
    "export_raid_plan_builds_xlsx",
    "raid_plan_build_export",
    "raid_plan_discord_builds_text",
]
